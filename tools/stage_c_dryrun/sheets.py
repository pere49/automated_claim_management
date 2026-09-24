"""PROTOTYPE (tools/) — Stage C dry run: claim items from a claim sheet.

Two sources, enough for the dry run (app/claims builds the real readers in C1):
  an Excel tab    the claim-form template's fixed layout, given in the pair's config
                  (date column, expense columns, first data row, Total label and column)
  an items file   rows already typed out (date cell + amounts), for a PDF sheet
Dates are read by the date trial's recommended rule over the whole sheet.
Not used by the application.
"""

from __future__ import annotations

import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

from claim_check import Item


def excel_sheet(path: Path, tab: str, layout: dict) -> tuple[list[tuple[int, object, list]], Decimal | None]:
    """([(row, date cell, [(column name, amount)])], grand Total)."""
    ws = openpyxl.load_workbook(path, data_only=True)[tab]
    first, last = (column_index_from_string(c) for c in layout["expense_columns"].split(":"))
    date_col = column_index_from_string(layout["date_column"])
    total_col = column_index_from_string(layout["total_column"])
    names = {c: " ".join(str(ws.cell(r, c).value or "").strip() for r in layout["header_rows"]).strip()
             for c in range(first, last + 1)}
    rows, grand = [], None
    for r in range(layout["first_row"], ws.max_row + 1):
        if any(str(ws.cell(r, c).value or "").strip() == layout["total_label"] for c in range(1, total_col + 1)):
            grand = Decimal(str(round(ws.cell(r, total_col).value, 2)))
            break
        amounts = [(names[c], Decimal(str(round(v, 2)))) for c in range(first, last + 1)
                   if isinstance((v := ws.cell(r, c).value), (int, float)) and v]
        rows.append((r, ws.cell(r, date_col).value, amounts))
    return rows, grand


def items_file(path: Path) -> tuple[list[tuple[int, object, list]], Decimal | None]:
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = [(i + 1, row["date"] or None, [(f"amount {k + 1}", Decimal(a)) for k, a in enumerate(row["amounts"])])
            for i, row in enumerate(data["rows"])]
    return rows, Decimal(data["grand_total"])


def claim_items(rows: list, read_dates) -> list[Item]:
    """read_dates: date cells of the whole sheet -> one date (or None) per cell."""
    dated = read_dates([cell for _, cell, _ in rows])
    return [Item(r, name, amount, d) for (r, _, amounts), d in zip(rows, dated) for name, amount in amounts]
