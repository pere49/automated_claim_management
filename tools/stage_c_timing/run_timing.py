"""PROTOTYPE (tools/) — Stage C: what each new step costs, next to OCR.

    .venv/Scripts/python.exe tools/stage_c_timing/run_timing.py

Times, on the real July receipts (their OCR readings, saved by the search
trial) and on larger made-up documents (copies of those pages with every
digit changed, so no two copies are alike):
  OCR                    measured per page when the readings were made
  claim sheet (Excel)    a made-up workbook the size of the sample, read with openpyxl
  claim-sheet dates      the R9 date rule (tools/date_parsing_trial) over the date cells
  page preparation       Stage B's once-per-page work, already paid as each page arrives
  cents index            per page, the whole parts of amounts printed with cents
  PIN scan               the PIN searched on every page
  repeat check           fingerprints, then text / moment / references (tools/duplicate_pages_trial),
                         pairwise and indexed
  check every amount     one claim per receipt page: amount (exact or cents dropped) + date
                         + PIN (required only when the PIN scan found it), each page used once;
                         two ways, which must give the same page for every claim:
                           every page    each claim searches every page (Stage B's search)
                           indexed       an index of every amount text printed (built once)
                                         picks the pages that carry the claim's amount; only
                                         those are searched in full
  date evidence          the two extra date searches one ambiguous claim date costs
Settings in timing_config.json. Claims come from the private answer key and
nothing from the documents is printed. Not used by the application.

Measured 2026-09-24 (26 pages = the July receipts; OCR 13.7 s a page, 6.0 min):
  Excel sheet 18-29 ms; its dates (R9, worst case) 11 ms; PIN scan 2 ms;
  repeat check indexed 3 ms (pairwise 8 ms; at 500 pages 75 ms vs 3.1 s);
  fingerprints 42 ms and cents index 2 ms (both can be made as pages arrive);
  check every amount: every page 347 ms, indexed 27 ms (100 pages: 5.8 s vs
  150 ms; 500 pages: 136 s vs 1.0 s), the same page for every claim.
"""

from __future__ import annotations

import json
import random
import re
import statistics
import sys
import tempfile
import time
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
for folder in (ROOT, ROOT / "tools/duplicate_pages_trial", ROOT / "tools/date_parsing_trial",
               ROOT / "tools/search_criteria_trial"):
    sys.path.insert(0, str(folder))

import openpyxl  # noqa: E402
from cases_real import parse_printed_date  # noqa: E402
from checks import Checks  # noqa: E402
from strategies import Parser, run_rule  # noqa: E402

from app.config_files import read_json_config  # noqa: E402
from app.layout import group_rows, load_row_rules  # noqa: E402
from app.matching import Query, load_matching_rules, prepare_page, search_document  # noqa: E402
from app.matching.amount_search import _faded, trim  # noqa: E402  (prototype: the app would export these)
from app.matching.money import amount_forms  # noqa: E402

CENTS = re.compile(r"(.+)[.,](\d{2})")


def timed(fn, repeats: int, slow: float) -> tuple[float, object]:
    """(median seconds, last result)."""
    times, result = [], None
    for _ in range(repeats):
        start = time.perf_counter()
        result = fn()
        times.append(time.perf_counter() - start)
        if times[-1] > slow:
            break
    return statistics.median(times), result


def load_document(cfg: dict) -> tuple[list[list[dict]], list[float], list[Query]]:
    """(each page's OCR words, each page's OCR seconds, one claim per page)."""
    readings = json.loads((ROOT / cfg["readings"]).read_text(encoding="utf-8"))["current"]
    key = json.loads((ROOT / cfg["answer_key"]).read_text(encoding="utf-8"))["pages"]
    prefix = cfg["document"] + "#"
    pages = sorted((p for p in readings if p["page"].startswith(prefix)), key=lambda p: int(p["page"].split("#")[1]))
    claims = []
    for p in pages:
        truth = key.get(p["page"], {})
        if truth.get("amounts") and truth.get("dates"):
            whole = Decimal(truth["amounts"][0]).to_integral_value(rounding="ROUND_FLOOR")  # typed without cents
            claims.append(Query(date=parse_printed_date(truth["dates"][0]), amount=whole, pin=cfg["fake_pin"]))
    return [p["words"] for p in pages], [p["seconds"] for p in pages], claims


def copies(pages: list[list[dict]], size: int, rnd: random.Random) -> list[list[dict]]:
    """`size` pages: the originals, then copies with every digit changed."""
    out = []
    for i in range(size):
        words = pages[i % len(pages)]
        if i >= len(pages):
            words = [dict(w, text=re.sub(r"\d", lambda _: str(rnd.randint(0, 9)), w["text"])) for w in words]
        out.append(words)
    return out


def cents_index(page) -> set[str]:
    return {m.group(1) for row in page.amount_cores for core, _ in row
            if (m := CENTS.fullmatch(core)) and m.group(2) != "00"}


def amount_index(pages: dict, floors: dict, rules) -> tuple[dict, dict]:
    """(amount text -> pages printing it, incl. faded-decimal joins; whole part of an amount with cents -> pages)."""
    printed, whole = defaultdict(set), defaultdict(set)
    for n, p in pages.items():
        for tokens, cores in zip(p.amounts, p.amount_cores):
            for core, _ in cores:
                printed[core].add(n)
            for _, text in _faded(tokens, rules):
                printed[trim(text, rules)[0]].add(n)
        for w in floors[n]:
            whole[w].add(n)
    return printed, whole


def check_all(pages: dict, floors: dict, claims: list[Query], rules, pin_required: bool,
              index: tuple[dict, dict] | None = None) -> list[int | None]:
    """Every claim, in order: first unused page with amount (exact, else cents
    dropped) + date (+ PIN). Returns each claim's page (None: not found)."""
    taken, assigned = set(), []
    for q in claims:
        forms = amount_forms(q.amount, rules)
        if index is None:
            searched = pages
        else:
            printed, whole = index
            wanted = set().union(*(printed.get(f, ()) for f in forms.with_cents | forms.without_cents),
                                 *(whole.get(w, ()) for w in forms.without_cents))
            searched = {n: pages[n] for n in sorted(wanted)}
        res = search_document(searched, q, rules)
        exact = [n for n in res.pages_with("amount") if n not in taken]
        dropped = [n for n in sorted(searched) if n not in taken and n not in exact and floors[n] & forms.without_cents]
        page = next((n for n in exact + dropped
                     if res.pages[n].found("date") and (res.pages[n].found("pin") or not pin_required)), None)
        if page is not None:
            taken.add(page)
        assigned.append(page)
    return assigned


def excel_read(cfg: dict, folder: Path) -> float:
    rnd = random.Random(cfg["seed"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([f"Column {c}" for c in range(cfg["sheet_columns"])])
    day = datetime(2026, 7, 1)
    for r in range(cfg["sheet_rows"]):
        ws.append([day + timedelta(days=r % 30), f"Item {r}"]
                  + [rnd.choice([None, rnd.randint(50, 9000), round(rnd.uniform(50, 9000), 2)])
                     for _ in range(cfg["sheet_columns"] - 2)])
    path = folder / "made_up_claim_sheet.xlsx"
    wb.save(path)

    def read():
        book = openpyxl.load_workbook(path, data_only=True)
        return [row for row in book.active.iter_rows(values_only=True)]
    seconds, _ = timed(read, cfg["repeats"], cfg["slow_step_seconds"])
    return seconds


def main() -> None:
    cfg = read_json_config(HERE / "timing_config.json", {"sizes": list, "repeats": int})
    dup_cfg = read_json_config(ROOT / "tools/duplicate_pages_trial/trial_config.json", {"invoice_labels": list})
    date_cfg = read_json_config(ROOT / "tools/date_parsing_trial/trial_config.json", {"month_names": dict})
    rules, row_rules = load_matching_rules(), load_row_rules()
    checks, parser = Checks(dup_cfg), Parser(date_cfg)
    repeats, slow = cfg["repeats"], cfg["slow_step_seconds"]
    words, ocr_seconds, base_claims = load_document(cfg)
    per_page_ocr = statistics.median(ocr_seconds)
    print(f"document: {len(words)} receipt pages, {len(base_claims)} claims from the answer key; "
          f"OCR measured {sum(ocr_seconds):.0f} s in all, median {per_page_ocr:.1f} s a page\n")

    with tempfile.TemporaryDirectory() as folder:
        excel = excel_read(cfg, Path(folder))
    rnd = random.Random(cfg["seed"])
    cells = [f"{rnd.randint(1, 12)}/{rnd.randint(1, 12)}/2026" for _ in range(cfg["sheet_rows"])]  # all ambiguous
    dates, _ = timed(lambda: run_rule("R9", parser, cells, [set()] * len(cells)), repeats, slow)
    print(f"claim sheet, Excel {cfg['sheet_rows']}x{cfg['sheet_columns']} cells: {1000 * excel:.1f} ms; "
          f"its dates by rule R9 (every cell ambiguous, the worst case): {1000 * dates:.2f} ms\n")

    table: dict[str, list[str]] = {}
    for size in cfg["sizes"]:
        doc = copies(words, size, random.Random(cfg["seed"] + size))
        claims = [base_claims[i % len(base_claims)] for i in range(size)]

        def prepare():
            out = {}
            for n, page_words in enumerate(doc, 1):
                grouped = group_rows([SimpleNamespace(text=w["text"], box=w["box"]) for w in page_words], row_rules)
                out[n] = (prepare_page(grouped, rules), [[s.text for s in r.segments] for r in grouped])
            return out
        t_prep, prepared = timed(prepare, 1, slow)
        pages = {n: p for n, (p, _) in prepared.items()}
        t_cents, floors = timed(lambda: {n: cents_index(p) for n, p in pages.items()}, repeats, slow)
        t_pin, pin_scan = timed(lambda: search_document(pages, Query(pin=cfg["fake_pin"]), rules), repeats, slow)
        pin_required = bool(pin_scan.pages_with("pin"))
        t_fp, fps = timed(lambda: [checks.fingerprint([w["text"] for w in doc[n - 1]], prepared[n][1])
                                   for n in sorted(prepared)], repeats, slow)
        t_pair, _ = timed(lambda: [checks.pairwise(fps, c) for c in ("text", "moment", "references")], repeats, slow)
        t_index, _ = timed(lambda: [checks.indexed(fps, c) for c in ("text", "moment", "references")], repeats, slow)
        t_all, every = timed(lambda: check_all(pages, floors, claims, rules, pin_required), repeats, slow)
        t_build, index = timed(lambda: amount_index(pages, floors, rules), repeats, slow)
        t_fast, fast = timed(lambda: check_all(pages, floors, claims, rules, pin_required, index), repeats, slow)
        assert every == fast, f"{size} pages: indexed check gave a different page for {sum(a != b for a, b in zip(every, fast))} claims"
        q = claims[0]
        t_evidence, _ = timed(lambda: [search_document(pages, Query(date=q.date + timedelta(days=d), amount=q.amount), rules)
                                       for d in (0, 1)], repeats, slow)
        ms = lambda s: f"{1000 * s:,.0f} ms" if s >= 0.01 else f"{1000 * s:.1f} ms"  # noqa: E731
        for label, value in (("OCR (measured, median page x pages)", f"{per_page_ocr * size / 60:,.1f} min"),
                             ("page preparation (Stage B, as pages arrive)", ms(t_prep)),
                             ("cents index", ms(t_cents)),
                             ("PIN scan", ms(t_pin)),
                             ("repeat check: fingerprints", ms(t_fp)),
                             ("repeat check: pairwise", ms(t_pair)),
                             ("repeat check: indexed", ms(t_index)),
                             ("check every amount: every page", ms(t_all)),
                             ("check every amount: indexed (build + check)", ms(t_build + t_fast)),
                             ("date evidence, one ambiguous date", ms(t_evidence))):
            table.setdefault(label, []).append(value)
        found = sum(a is not None for a in every)
        print(f"  {size} pages ({size} claims) timed; PIN required: {pin_required}; claims given a page: {found}; "
              f"indexed check gave the same page for every claim")

    print(f"\n{'step':46}" + "".join(f"{str(s) + ' pages':>14}" for s in cfg["sizes"]))
    for label, values in table.items():
        print(f"{label:46}" + "".join(f"{v:>14}" for v in values))


if __name__ == "__main__":
    main()
